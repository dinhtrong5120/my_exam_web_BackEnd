from io import BytesIO
import re
from django.contrib.auth import authenticate, get_user_model
from django.db import transaction
from django.db.models import Count, Q
from django.http import Http404
from django.shortcuts import get_object_or_404
from rest_framework.authtoken.models import Token
from rest_framework.decorators import api_view, permission_classes, authentication_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from openpyxl import load_workbook
from .models import Exam, Question, Choice, Attempt

def error(message, code=400):
    return Response({'detail': message}, status=code)

def exam_list_data(exam):
    return {'id': exam.id, 'title': exam.title, 'description': exam.description,
            'category': exam.category, 'duration_minutes': exam.duration_minutes,
            'is_published': exam.is_published, 'question_count': exam.question_count,
            'attempt_count': exam.attempt_count, 'owner': exam.owner.username,
            'updated_at': exam.updated_at.isoformat()}

def exam_data(exam, reveal_answers=False):
    data = {'id': exam.id, 'title': exam.title, 'description': exam.description,
            'category': exam.category, 'duration_minutes': exam.duration_minutes,
            'is_published': exam.is_published, 'owner': exam.owner.username,
            'questions': []}
    for q in exam.questions.prefetch_related('choices').all():
        item = {'id': q.id, 'position': q.position, 'part': q.part, 'prompt': q.prompt,
                'question_type': q.question_type,
                'passage': q.passage, 'points': q.points, 'difficulty': q.difficulty,
                'image_url': q.image.url if q.image else '', 'audio_url': q.audio.url if q.audio else '',
                'choices': [{'id': c.id, 'text': c.text, 'position': c.position,
                             **({'is_correct': c.is_correct} if reveal_answers else {})} for c in q.choices.all()]}
        if reveal_answers:
            item['explanation'] = q.explanation
        data['questions'].append(item)
    return data

def own_exam(request, exam_id):
    return get_object_or_404(Exam, pk=exam_id, owner=request.user)

def validate_questions(items):
    if not isinstance(items, list):
        return 'Danh sách câu hỏi không hợp lệ.'
    for n, q in enumerate(items, 1):
        if not str(q.get('prompt', '')).strip():
            return f'Câu {n} cần có nội dung.'
        choices = q.get('choices', [])
        if not isinstance(choices, list) or len(choices) < 2 or any(not str(c.get('text', '')).strip() for c in choices):
            return f'Câu {n} cần ít nhất 2 phương án có nội dung.'
        kind = q.get('question_type') or 'single'
        if kind not in ('single', 'multiple'):
            return f'Câu {n} có loại câu hỏi không hợp lệ.'
        correct_count = sum(bool(c.get('is_correct')) for c in choices)
        if kind == 'single' and correct_count != 1:
            return f'Câu {n} chọn một đáp án phải có đúng một đáp án đúng.'
        if kind == 'multiple' and correct_count < 2:
            return f'Câu {n} chọn nhiều đáp án phải có ít nhất hai đáp án đúng.'
        try:
            if int(q.get('points', 1)) < 1:
                return f'Điểm câu {n} phải lớn hơn 0.'
        except (TypeError, ValueError):
            return f'Điểm câu {n} không hợp lệ.'
    return None

def save_questions(exam, items):
    # Existing question IDs retain uploaded media when the editor saves.
    existing = {q.id: q for q in exam.questions.all()}
    retained = set()
    for position, item in enumerate(items, 1):
        qid = item.get('id')
        q = existing.get(int(qid)) if str(qid).isdigit() else None
        if q is None:
            q = Question(exam=exam)
        q.position = position
        q.part = str(item.get('part') or 'Part 5')[:30]
        q.question_type = str(item.get('question_type') or 'single')
        q.prompt = str(item['prompt']).strip()
        q.passage = str(item.get('passage') or '')
        q.explanation = str(item.get('explanation') or '')
        q.points = int(item.get('points', 1))
        q.difficulty = str(item.get('difficulty') or 'easy')[:20]
        q.save()
        retained.add(q.id)
        q.choices.all().delete()
        Choice.objects.bulk_create([Choice(question=q, text=str(c['text']).strip(),
            is_correct=bool(c['is_correct']), position=i) for i, c in enumerate(item['choices'])])
    for q in exam.questions.exclude(id__in=retained):
        q.delete()

@api_view(['POST'])
@permission_classes([AllowAny])
def register(request):
    username = str(request.data.get('username', '')).strip()
    password = str(request.data.get('password', ''))
    if len(username) < 3 or len(password) < 8:
        return error('Tên đăng nhập cần từ 3 ký tự, mật khẩu từ 8 ký tự.')
    User = get_user_model()
    if User.objects.filter(username__iexact=username).exists():
        return error('Tên đăng nhập đã tồn tại.')
    user = User.objects.create_user(username=username, password=password)
    token, _ = Token.objects.get_or_create(user=user)
    return Response({'token': token.key, 'username': user.username}, status=201)

@api_view(['POST'])
@permission_classes([AllowAny])
def login(request):
    user = authenticate(username=request.data.get('username'), password=request.data.get('password'))
    if user is None:
        return error('Tên đăng nhập hoặc mật khẩu không đúng.', 401)
    token, _ = Token.objects.get_or_create(user=user)
    return Response({'token': token.key, 'username': user.username})

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def me(request):
    return Response({'username': request.user.username, 'id': request.user.id})

@api_view(['GET', 'POST'])
def exams(request):
    if request.method == 'GET':
        qs = Exam.objects.select_related('owner').annotate(question_count=Count('questions', distinct=True),
            attempt_count=Count('attempts', distinct=True))
        if request.query_params.get('mine') == '1':
            if not request.user.is_authenticated:
                return error('Vui lòng đăng nhập.', 401)
            qs = qs.filter(owner=request.user)
        else:
            qs = qs.filter(is_published=True)
        search = request.query_params.get('search', '').strip()
        if search:
            qs = qs.filter(Q(title__icontains=search) | Q(description__icontains=search))
        return Response([exam_list_data(e) for e in qs[:100]])
    if not request.user.is_authenticated:
        return error('Vui lòng đăng nhập.', 401)
    title = str(request.data.get('title', '')).strip()
    if not title:
        return error('Vui lòng nhập tên bài thi.')
    try:
        duration = int(request.data.get('duration_minutes', 30))
    except (TypeError, ValueError):
        return error('Thời gian làm bài không hợp lệ.')
    if not 1 <= duration <= 300:
        return error('Thời gian làm bài phải từ 1 đến 300 phút.')
    exam = Exam.objects.create(owner=request.user, title=title[:200],
        description=str(request.data.get('description', '')),
        category=str(request.data.get('category') or 'TOEIC')[:60], duration_minutes=duration)
    return Response(exam_data(exam, True), status=201)

@api_view(['GET', 'PATCH', 'DELETE'])
def exam_detail(request, exam_id):
    exam = get_object_or_404(Exam.objects.select_related('owner'), pk=exam_id)
    owns = request.user.is_authenticated and request.user.id == exam.owner_id
    if request.method == 'GET':
        if not exam.is_published and not owns:
            raise Http404
        return Response(exam_data(exam, owns))
    if not owns:
        return error('Bạn không có quyền sửa bài thi này.', 403)
    if request.method == 'DELETE':
        with transaction.atomic():
            exam.attempts.update(exam_title=exam.title,
                exam_description=exam.description, exam_category=exam.category)
            exam.delete()
        return Response(status=204)
    items = request.data.get('questions')
    if items is not None:
        problem = validate_questions(items)
        if problem:
            return error(problem)
    try:
        duration = int(request.data.get('duration_minutes', exam.duration_minutes))
    except (TypeError, ValueError):
        return error('Thời gian làm bài không hợp lệ.')
    if not 1 <= duration <= 300:
        return error('Thời gian làm bài phải từ 1 đến 300 phút.')
    title = str(request.data.get('title', exam.title)).strip()
    if not title:
        return error('Vui lòng nhập tên bài thi.')
    with transaction.atomic():
        exam.title = title[:200]
        exam.description = str(request.data.get('description', exam.description))
        exam.category = str(request.data.get('category', exam.category))[:60]
        exam.duration_minutes = duration
        exam.save()
        if items is not None:
            save_questions(exam, items)
    return Response(exam_data(exam, True))

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def publish_exam(request, exam_id):
    exam = own_exam(request, exam_id)
    if not exam.questions.exists():
        return error('Thêm ít nhất một câu hỏi trước khi xuất bản.')
    exam.is_published = bool(request.data.get('is_published', True))
    exam.save(update_fields=['is_published', 'updated_at'])
    return Response({'is_published': exam.is_published})

@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def import_questions(request, exam_id):
    exam = own_exam(request, exam_id)
    file = request.FILES.get('file')
    if not file or not file.name.lower().endswith('.xlsx') or file.size > 5 * 1024 * 1024:
        return error('Chọn tệp .xlsx tối đa 5 MB.')
    try:
        sheet = load_workbook(BytesIO(file.read()), read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        header = [str(v or '').strip() for v in next(rows)]
        required = ['Nội dung câu hỏi', 'Phương án 1', 'Đúng/Sai 1', 'Phương án 2', 'Đúng/Sai 2']
        if any(x not in header for x in required):
            return error('Tệp không đúng mẫu cột câu hỏi. Dùng mẫu trong thư mục Template.')
        cols = {name: i for i, name in enumerate(header)}
        option_numbers = sorted(
            int(match.group(1))
            for name in header
            if (match := re.fullmatch(r'Phương án (\d+)', name))
        )
        parsed = []
        for line, row in enumerate(rows, 2):
            get = lambda name, default=None: row[cols[name]] if name in cols and cols[name] < len(row) else default
            prompt = get('Nội dung câu hỏi')
            if not prompt:
                continue
            choices = []
            for n in option_numbers:
                value = get(f'Phương án {n}')
                if value is not None and str(value).strip():
                    correct = str(get(f'Đúng/Sai {n}', '')).strip().lower() in ('true', '1', 'yes', 'đúng')
                    choices.append({'text': str(value), 'is_correct': correct})
            item = {'prompt': str(prompt), 'explanation': str(get('Giải thích') or ''),
                    'points': get('Điểm') or 1, 'difficulty': str(get('Độ khó') or 'easy'),
                    'part': str(get('Part') or 'Part 5'),
                    'question_type': str(get('Loại câu hỏi') or 'single').strip().lower(),
                    'choices': choices}
            problem = validate_questions([item])
            if problem:
                return error(f'Dòng {line}: {problem}')
            parsed.append(item)
            if len(parsed) > 500:
                return error('Tối đa 500 câu hỏi mỗi lần nhập.')
        if not parsed:
            return error('Không tìm thấy câu hỏi trong tệp.')
        with transaction.atomic():
            start = exam.questions.count()
            for i, item in enumerate(parsed, start + 1):
                q = Question.objects.create(exam=exam, position=i, part=item['part'][:30],
                    question_type=item['question_type'],
                    prompt=item['prompt'], explanation=item['explanation'],
                    points=int(item['points']), difficulty=item['difficulty'][:20])
                Choice.objects.bulk_create([Choice(question=q, text=c['text'], is_correct=c['is_correct'],
                    position=n) for n, c in enumerate(item['choices'])])
        return Response({'imported': len(parsed), 'exam': exam_data(exam, True)})
    except (ValueError, KeyError, StopIteration, OSError, TypeError) as exc:
        return error(f'Không đọc được tệp Excel: {exc}')

@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def upload_question_media(request, exam_id, question_id):
    exam = own_exam(request, exam_id)
    q = get_object_or_404(Question, pk=question_id, exam=exam)
    kind = request.data.get('kind')
    file = request.FILES.get('file')
    if kind not in ('image', 'audio') or not file:
        return error('Chọn ảnh hoặc âm thanh để tải lên.')
    accepted = {'image': ('image/jpeg', 'image/png', 'image/webp'),
                'audio': ('audio/mpeg', 'audio/wav', 'audio/ogg', 'audio/mp4')}
    if file.content_type not in accepted[kind] or file.size > 10 * 1024 * 1024:
        return error('Định dạng tệp không hỗ trợ hoặc vượt quá 10 MB.')
    setattr(q, kind, file)
    q.save(update_fields=[kind])
    return Response({'url': getattr(q, kind).url})

@api_view(['POST'])
@permission_classes([AllowAny])
def submit_exam(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id, is_published=True)
    answers = request.data.get('answers', {})
    if not isinstance(answers, dict):
        return error('Đáp án không hợp lệ.')
    questions = list(exam.questions.prefetch_related('choices'))
    score = 0
    maximum = sum(q.points for q in questions)
    details = []
    review_questions = []
    for q in questions:
        correct_ids = [c.id for c in q.choices.all() if c.is_correct]
        selected = answers.get(str(q.id))
        if q.question_type == 'multiple':
            selected_ids = selected if isinstance(selected, list) else []
            right = {str(x) for x in selected_ids} == {str(x) for x in correct_ids}
        else:
            selected_ids = [selected] if selected is not None else []
            right = len(selected_ids) == 1 and str(selected_ids[0]) == str(correct_ids[0]) if correct_ids else False
        score += q.points if right else 0
        details.append({'question_id': q.id, 'selected_choice_ids': selected_ids,
            'correct_choice_ids': correct_ids, 'is_correct': right,
            'explanation': q.explanation})
        review_questions.append({'id': q.id, 'prompt': q.prompt, 'part': q.part,
            'passage': q.passage, 'question_type': q.question_type, 'points': q.points,
            'image_url': q.image.url if q.image else '',
            'audio_url': q.audio.url if q.audio else '',
            'choices': [{'id': c.id, 'text': c.text} for c in q.choices.all()],
            'selected_choice_ids': selected_ids, 'correct_choice_ids': correct_ids,
            'is_correct': right, 'explanation': q.explanation})
    attempt = Attempt.objects.create(exam=exam,
        exam_title=exam.title, exam_description=exam.description,
        exam_category=exam.category,
        user=request.user if request.user.is_authenticated else None,
        guest_name=str(request.data.get('guest_name', 'Khách'))[:80],
        answers=answers, score=score, max_score=maximum,
        result_snapshot={'questions': review_questions})
    return Response({'id': attempt.id, 'score': score, 'max_score': maximum,
        'correct_count': sum(d['is_correct'] for d in details), 'total': len(questions), 'details': details}, status=201)

def attempt_data(attempt):
    questions = attempt.result_snapshot.get('questions', [])
    return {'id': attempt.id, 'exam_id': attempt.exam_id,
        'exam_title': attempt.exam.title if attempt.exam else attempt.exam_title or 'Bài thi đã xóa',
        'exam_description': attempt.exam.description if attempt.exam else attempt.exam_description,
        'exam_category': attempt.exam.category if attempt.exam else attempt.exam_category,
        'score': attempt.score, 'max_score': attempt.max_score,
        'correct_count': sum(q['is_correct'] for q in questions) if questions else None,
        'total': len(questions) if questions else None,
        'submitted_at': attempt.submitted_at.isoformat()}

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def my_attempts(request):
    attempts = Attempt.objects.filter(user=request.user).select_related('exam')[:100]
    return Response([attempt_data(attempt) for attempt in attempts])

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def my_attempt_detail(request, attempt_id):
    attempt = get_object_or_404(Attempt.objects.select_related('exam'), pk=attempt_id, user=request.user)
    return Response({**attempt_data(attempt),
        'questions': attempt.result_snapshot.get('questions', []),
        'has_snapshot': bool(attempt.result_snapshot.get('questions'))})

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def exam_attempts(request, exam_id):
    exam = own_exam(request, exam_id)
    return Response([{'id': a.id, 'name': a.user.username if a.user else a.guest_name,
        'score': a.score, 'max_score': a.max_score, 'submitted_at': a.submitted_at.isoformat()}
        for a in exam.attempts.select_related('user')[:100]])
